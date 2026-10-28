"""Smoke/unit tests against a fake Cloudflare API (httpx MockTransport)."""
import io, json, os, sys

os.environ.setdefault("TESTDATA", "/tmp/cfzt-test-data")
os.environ.update({
    "CF_API_TOKEN": "fake-token", "CF_ACCOUNT_ID": "acct123",
    "CF0T_PASSWORD": "secret", "CF0T_SECRET_KEY": "x" * 32,
    "CF0T_DATA_DIR": os.environ["TESTDATA"],
})
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import httpx
import pytest
from fastapi.testclient import TestClient
from app import auth, backups, cf, changelog, services, main
from app.cf import CloudflareError
from app.config import settings

TUNNEL = "tunnel-abc-123"
VALID_TOKENS = {"fake-token", "replacement-token"}
STATE = {}


def reset_state():
    STATE.clear()
    STATE.update({
        "config": {"ingress": [
            {"hostname": "ha.makooy.nl", "service": "https://192.168.1.12:8123",
             "originRequest": {"noTLSVerify": True}},
            {"hostname": "sonarr.makooy.nl", "service": "http://192.168.1.9:8989"},
            {"service": "http_status:404"},
        ]},
        "apps": [{"id": "app1", "name": "Sonarr", "domain": "sonarr.makooy.nl",
                  "type": "self_hosted", "session_duration": "24h", "aud": "aud-hash",
                  "created_at": "2026-01-01T00:00:00Z",
                  "allowed_idps": ["google-idp-id"],
                  "auto_redirect_to_identity": True,
                  "custom_deny_message": "Ask Marc for access",
                  "enable_binding_cookie": True,
                  "destinations": [{"type": "public", "uri": "sonarr.makooy.nl"}],
                  "policies": [{"id": "pol1", "name": "GoogleIDP-Marc-Policy", "precedence": 1}]},
                 {"id": "app2", "name": "Google Workspace", "domain": "workspace.makooy.nl",
                  "type": "saas", "policies": []},
                 {"id": "app3", "name": "Books", "domain": "readarr.makooy.nl",
                  "type": "self_hosted", "policies": [],
                  "destinations": [{"type": "public", "uri": "readarr.makooy.nl"},
                                   {"type": "public", "uri": "calibre.makooy.nl"}]}],
        "dns": [{"id": "dns1", "type": "CNAME", "name": "sonarr.makooy.nl",
                 "content": f"{TUNNEL}.cfargotunnel.com", "proxied": True}],
    })


reset_state()


def handler(request: httpx.Request) -> httpx.Response:
    path, method = request.url.path, request.method
    body = json.loads(request.content) if request.content else None
    page = request.url.params.get("page", "1")

    def ok(result):
        return httpx.Response(200, json={"success": True, "errors": [], "result": result})

    token = request.headers.get("Authorization", "").removeprefix("Bearer ")
    if token not in VALID_TOKENS:
        return httpx.Response(400, json={"success": False, "result": None, "errors": [
            {"code": 6003, "message": "Invalid request headers"}]})

    if path.endswith("/user/tokens/verify"):
        return ok({"status": "active"})
    if path.endswith("/cfd_tunnel") and method == "GET":
        return ok([{"id": TUNNEL, "name": "HATunnel", "status": "healthy"}] if page == "1" else [])
    if path.endswith(f"/cfd_tunnel/{TUNNEL}"):
        return ok({"id": TUNNEL, "name": "HATunnel", "status": "healthy"})
    if path.endswith("/configurations"):
        if method == "PUT":
            STATE["config"] = body["config"]
            return ok({"config": body["config"]})
        return ok({"config": STATE["config"]})
    if path.endswith("/access/apps") and method == "GET":
        return ok(STATE["apps"] if page == "1" else [])
    if path.endswith("/access/apps") and method == "POST":
        new = dict(body, id=f"app{len(STATE['apps']) + 1}")
        new["policies"] = [{"id": p["id"], "name": "pol", "precedence": p["precedence"]}
                           for p in body.get("policies", [])]
        STATE["apps"].append(new)
        return ok(new)
    if "/access/apps/" in path and method == "PUT":
        aid = path.rsplit("/", 1)[-1]
        for field in ("id", "aud", "created_at", "updated_at"):
            assert field not in body, f"{field} is read-only and must not be sent back"
        # Cloudflare replaces the whole object: whatever is not sent is lost.
        for i, a in enumerate(STATE["apps"]):
            if a["id"] == aid:
                STATE["apps"][i] = {"id": a["id"], "aud": a.get("aud"),
                                    "created_at": a.get("created_at"), **body}
        return ok(body)
    if "/access/apps/" in path and method == "DELETE":
        aid = path.rsplit("/", 1)[-1]
        STATE["apps"] = [a for a in STATE["apps"] if a["id"] != aid]
        return ok({"id": aid})
    if path.endswith("/access/identity_providers"):
        return ok([{"id": "idp-google", "name": "Google", "type": "google"},
                   {"id": "idp-otp", "name": "", "type": "onetimepin"}]
                  if page == "1" else [])
    if path.endswith("/access/policies"):
        return ok([{"id": "pol1", "name": "GoogleIDP-Marc-Policy", "decision": "allow"},
                   {"id": "pol2", "name": "Everyone-Policy", "decision": "allow"}] if page == "1" else [])
    if path == "/client/v4/zones":
        return ok([{"id": "zone1", "name": "makooy.nl"}] if page == "1" else [])
    if path.endswith("/dns_records") and method == "GET":
        name = request.url.params.get("name")
        recs = [r for r in STATE["dns"] if not name or r["name"] == name]
        return ok(recs if page == "1" else [])
    if path.endswith("/dns_records") and method == "POST":
        rec = dict(body, id=f"dns{len(STATE['dns']) + 1}")
        STATE["dns"].append(rec)
        return ok(rec)
    if "/dns_records/" in path and method == "PATCH":
        rid = path.rsplit("/", 1)[-1]
        for i, r in enumerate(STATE["dns"]):
            if r["id"] == rid:
                STATE["dns"][i] = dict(r, **body)
                return ok(STATE["dns"][i])
        return ok({"id": rid})
    if "/dns_records/" in path and method == "DELETE":
        rid = path.rsplit("/", 1)[-1]
        STATE["dns"] = [r for r in STATE["dns"] if r["id"] != rid]
        return ok({"id": rid})
    return httpx.Response(404, json={"success": False, "result": None, "errors": [
        {"code": 1, "message": f"no mock for {method} {path}"}]})


def make_client(token: str, account_id: str, timeout: float = 30.0) -> cf.CloudflareClient:
    """A CloudflareClient whose transport is the mock above."""
    client = cf.CloudflareClient.__new__(cf.CloudflareClient)
    client.account_id = account_id
    client._client = httpx.AsyncClient(
        base_url=cf.API_BASE, transport=httpx.MockTransport(handler),
        headers={"Authorization": f"Bearer {token}"},
    )
    return client


main.CloudflareClient = make_client
settings.store.clear()
settings.save(tunnel_id=TUNNEL)

tc = TestClient(main.app)


@pytest.fixture(autouse=True)
def clean_state():
    main.reset_client()
    services.invalidate_overview()
    yield


def post(client, url, data=None, **kwargs):
    """POST the way a browser does: echo the CSRF cookie back in the form."""
    token = client.cookies.get(main.CSRF_COOKIE)
    if not token:
        client.get("/login")
        token = client.cookies.get(main.CSRF_COOKIE)
    payload = dict(data or {})
    payload.setdefault("csrf_token", token)
    return client.post(url, data=payload, **kwargs)


def login(client=tc, password="secret"):
    r = post(client, "/login", data={"password": password}, follow_redirects=False)
    assert r.status_code == 303, r.text
    return r


# ----------------------------------------------------------------- basics ---
def test_auth_required():
    auth.reset_throttle()
    fresh = TestClient(main.app)
    r = fresh.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    r = post(fresh, "/login", data={"password": "wrong"}, follow_redirects=False)
    assert "err=" in r.headers["location"]


def test_a_form_without_a_csrf_token_is_refused():
    auth.reset_throttle()
    fresh = TestClient(main.app)
    fresh.get("/login")  # gets the cookie, which we then do not echo back
    r = fresh.post("/login", data={"password": "secret"}, follow_redirects=False)
    assert r.status_code == 400, "a POST that does not carry the token is rejected"
    assert "expired" in r.text

    r = fresh.post("/login", data={"password": "secret", "csrf_token": "something else"},
                   follow_redirects=False)
    assert r.status_code == 400, "and a wrong token is rejected too"

    r = post(fresh, "/login", data={"password": "secret"}, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers.get("location", "err=")


def test_login_is_throttled():
    auth.reset_throttle()
    fresh = TestClient(main.app)
    for attempt in range(auth.MAX_ATTEMPTS):
        r = post(fresh, "/login", data={"password": "wrong"}, follow_redirects=False)
        assert "err=" in r.headers["location"]
    # the next one is refused before the password is even looked at
    r = post(fresh, "/login", data={"password": "secret"}, follow_redirects=False)
    assert "Too+many+attempts" in r.headers["location"]
    assert auth.lockout_remaining("testclient"), "the address is locked out"

    auth.reset_throttle()
    r = post(fresh, "/login", data={"password": "secret"}, follow_redirects=False)
    assert "err=" not in r.headers.get("location", "err="), "and can log in once cleared"


def test_a_failed_attempt_does_not_block_the_event_loop():
    """The delay after a wrong password must be awaited, not slept through:
    time.sleep in an async route freezes every other request."""
    import inspect
    code = [line.split("#")[0] for line in inspect.getsource(main.login).splitlines()]
    assert any("await asyncio.sleep" in line for line in code)
    assert not any("time.sleep" in line for line in code)


def test_dashboard():
    reset_state()
    login()
    r = tc.get("/")
    assert r.status_code == 200
    assert "ha.makooy.nl" in r.text and "sonarr.makooy.nl" in r.text
    assert "GoogleIDP-Marc-Policy" in r.text
    assert "https://192.168.1.12:8123" in r.text


# ------------------------------------------------------------------- crud ---
def test_create_full_app():
    reset_state()
    login()
    r = post(tc, "/new", data={
        "hostname": "jellyfin", "zone": "makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8096",
        "create_app": "1", "app_name": "Jellyfin", "session_duration": "24h",
        "allowed_idps": ["idp-google"],
        "policies": ["pol1", "pol2"], "launcher_visible": "1", "create_dns": "1",
    }, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]

    rule = cf.find_rule(STATE["config"]["ingress"], "jellyfin.makooy.nl")
    assert rule and rule["service"] == "http://192.168.1.9:8096"
    assert cf.is_catch_all(STATE["config"]["ingress"][-1]), "catch-all must stay last"
    app = [a for a in STATE["apps"] if a["domain"] == "jellyfin.makooy.nl"][0]
    assert app["type"] == "self_hosted"
    assert [p["id"] for p in app["policies"]] == ["pol1", "pol2"], "both policies attached"
    record = [d for d in STATE["dns"] if d["name"] == "jellyfin.makooy.nl"][0]
    assert record["content"].endswith("cfargotunnel.com")
    assert "comment" not in record, "do not write anything of our own into the user's DNS record"


def test_edit_form_splits_the_hostname_from_the_domain():
    """The hostname field and the domain dropdown must not repeat the zone."""
    reset_state()
    login()
    page = tc.get("/edit?h=sonarr.makooy.nl&p=").text
    assert 'name="hostname" value="sonarr"' in page, "only the subdomain belongs in the field"
    assert 'value="makooy.nl" selected' in page, "the zone belongs in the dropdown"
    assert 'value="sonarr.makooy.nl"' not in page.split('name="original_hostname"')[0]


def test_hostname_splitting():
    zones = [{"id": "z", "name": "makooy.nl"}]
    assert services.split_hostname("sonarr.makooy.nl", zones) == ("sonarr", "makooy.nl")
    assert services.split_hostname("a.b.makooy.nl", zones) == ("a.b", "makooy.nl")
    assert services.split_hostname("makooy.nl", zones) == ("makooy.nl", ""), "apex has no subdomain"
    assert services.split_hostname("host.example.org", zones) == ("host.example.org", "")
    assert services.split_hostname("host.example.org", []) == ("host.example.org", "")


def test_edit_changes_service_and_policies():
    reset_state()
    login()
    r = tc.get("/edit?h=sonarr.makooy.nl&p=")
    assert r.status_code == 200 and "8989" in r.text
    r = post(tc, "/edit", data={
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "sonarr.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "9999", "noTLSVerify": "1",
        "create_app": "1", "app_name": "Sonarr", "session_duration": "6h",
        "allowed_idps": ["idp-google"],
        "policies": ["pol2"], "launcher_visible": "1",
    }, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    rule = cf.find_rule(STATE["config"]["ingress"], "sonarr.makooy.nl")
    assert rule["service"] == "http://192.168.1.9:9999"
    assert rule["originRequest"]["noTLSVerify"] is True
    app = [a for a in STATE["apps"] if a["domain"] == "sonarr.makooy.nl"][0]
    assert app["session_duration"] == "6h"
    assert len(STATE["config"]["ingress"]) == 3, "no duplicate rule"


def test_rename_carries_the_dns_record_over():
    """Renaming a hostname must not leave the new one without DNS (error 1016)."""
    reset_state()
    login()
    # sonarr has a dedicated CNAME; the DNS checkbox is deliberately NOT sent,
    # which is what the edit form does for a hostname that already has one.
    r = post(tc, "/edit", data={
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "series.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8989",
        "create_app": "1", "app_name": "Sonarr", "session_duration": "24h",
        "allowed_idps": ["idp-google"],
        "policies": ["pol1"], "launcher_visible": "1",
    }, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]

    new = [d for d in STATE["dns"] if d["name"] == "series.makooy.nl"]
    assert new and new[0]["content"].endswith(".cfargotunnel.com"), "new hostname needs DNS"
    assert cf.find_rule(STATE["config"]["ingress"], "series.makooy.nl")
    assert cf.find_rule(STATE["config"]["ingress"], "sonarr.makooy.nl") is None


def test_leftover_dns_is_listed_so_it_can_be_removed():
    reset_state()
    # A record pointing at the tunnel with no route and no Access app behind it.
    STATE["config"] = cf.remove_rule(STATE["config"], "sonarr.makooy.nl")
    STATE["apps"] = []
    login()
    page = tc.get("/").text
    assert "sonarr.makooy.nl" in page and "unused DNS" in page

    r = post(tc, "/delete", data={"hostname": "sonarr.makooy.nl", "path": "", "del_dns_ids": "dns1"},
                follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    assert not [d for d in STATE["dns"] if d["name"] == "sonarr.makooy.nl"]


def test_edit_keeps_settings_this_tool_does_not_manage():
    """A PUT replaces the whole app, so identity providers and friends must be
    carried over instead of silently falling back to the defaults."""
    reset_state()
    login()
    r = post(tc, "/edit", data={
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "series.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8989",
        "create_app": "1", "app_name": "Series", "session_duration": "6h",
        "allowed_idps": ["idp-google"],
        "policies": ["pol2"], "launcher_visible": "1",
    }, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]

    app = [a for a in STATE["apps"] if a["id"] == "app1"][0]
    # Login methods and instant auth are the form's own fields now, so they are
    # whatever was submitted. Everything the form does not know about must
    # still come through untouched.
    assert app["allowed_idps"] == ["idp-google"]
    assert app["custom_deny_message"] == "Ask Marc for access"
    assert app["enable_binding_cookie"] is True
    assert app["destinations"] == [{"type": "public", "uri": "series.makooy.nl"}], \
        "a rename moves the hostname in destinations too"
    # ...while the fields the tool does manage are updated
    assert app["name"] == "Series" and app["session_duration"] == "6h"
    assert [p["id"] for p in app["policies"]] == ["pol2"]
    assert app["domain"] == "series.makooy.nl"


def test_adding_a_hostname_that_already_has_a_route_is_refused():
    reset_state()
    login()
    before = list(STATE["config"]["ingress"])
    r = post(tc, "/new", data={
        "hostname": "sonarr", "zone": "makooy.nl", "scheme": "http",
        "origin_host": "10.0.0.1", "origin_port": "1234",
    }, follow_redirects=False)
    assert "err=" in r.headers["location"], "must not silently replace the existing route"
    assert "already+has+a+tunnel+route" in r.headers["location"]
    assert STATE["config"]["ingress"] == before, "nothing was changed"


def test_renaming_onto_an_existing_hostname_is_refused():
    reset_state()
    login()
    r = post(tc, "/edit", data={
        "original_hostname": "ha.makooy.nl", "original_path": "",
        "hostname": "sonarr.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.12", "origin_port": "8123",
    }, follow_redirects=False)
    assert "err=" in r.headers["location"]
    assert cf.find_rule(STATE["config"]["ingress"], "ha.makooy.nl"), "original left alone"
    assert len([r for r in STATE["config"]["ingress"]
                if r.get("hostname") == "sonarr.makooy.nl"]) == 1, "no duplicate rule"


def test_edit_and_delete_target_the_exact_path():
    """A hostname with several paths must not fall back to the wrong rule."""
    reset_state()
    STATE["config"]["ingress"].insert(1, {"hostname": "ha.makooy.nl", "path": "/api",
                                          "service": "http://192.168.1.12:9000"})
    login()
    r = tc.get("/edit?h=ha.makooy.nl&p=/api")
    assert "9000" in r.text and "8123" not in r.text, "the /api rule is the one being edited"

    r = post(tc, "/delete", data={"hostname": "ha.makooy.nl", "path": "/api", "del_route": "1"},
                follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    assert cf.find_rule(STATE["config"]["ingress"], "ha.makooy.nl", "/api") is None
    assert cf.find_rule(STATE["config"]["ingress"], "ha.makooy.nl"), "the catch-all rule survived"

    # A path that does not exist must not hit a different rule.
    r = tc.get("/delete?h=ha.makooy.nl&p=/nope", follow_redirects=False)
    assert r.status_code == 303 and "not+found" in r.headers["location"]


def test_restore_keeps_every_setting_in_the_backup():
    reset_state()
    login()
    data = tc.get("/backup/export").json()
    STATE["apps"] = []
    r = post(tc, "/backup/preview", files={"backup": ("b.json", json.dumps(data), "application/json")})
    token = r.text.split('name="token" value="')[1].split('"')[0]
    r = post(tc, "/backup/apply", data={"token": token, "restore_apps": "1"},
                follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]

    restored = [a for a in STATE["apps"] if a["domain"] == "sonarr.makooy.nl"][0]
    assert restored["allowed_idps"] == ["google-idp-id"]
    assert restored["auto_redirect_to_identity"] is True
    assert restored["custom_deny_message"] == "Ask Marc for access"
    assert [p["id"] for p in restored["policies"]] == ["pol1"]


def test_deleting_an_app_asks_first():
    reset_state()
    login()
    data = {
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "sonarr.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8989",
        "policies": ["pol1", "pol2"], "noTLSVerify": "1",
    }  # note: no create_app, so the Access app would be deleted
    r = post(tc, "/edit", data=data)
    assert r.status_code == 200 and "Are you sure?" in r.text
    assert "without any authentication" in r.text
    # the whole submission is carried through, including multi-value fields
    assert 'name="policies" value="pol1"' in r.text
    assert 'name="policies" value="pol2"' in r.text
    assert 'name="noTLSVerify" value="1"' in r.text
    assert [a for a in STATE["apps"] if a["id"] == "app1"], "app still there"

    r = post(tc, "/edit", data={**data, "confirmed": "1"}, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    assert not [a for a in STATE["apps"] if a["id"] == "app1"], "deleted after confirming"


def test_removing_every_policy_asks_first():
    reset_state()
    login()
    data = {
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "sonarr.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8989",
        "create_app": "1", "app_name": "Sonarr", "session_duration": "24h",
        "allowed_idps": ["idp-google"],
    }  # no policies selected
    r = post(tc, "/edit", data=data)
    assert r.status_code == 200 and "Are you sure?" in r.text
    assert "denies everyone" in r.text
    app = [a for a in STATE["apps"] if a["id"] == "app1"][0]
    assert [p["id"] for p in app["policies"]] == ["pol1"], "policies untouched"

    r = post(tc, "/edit", data={**data, "confirmed": "1"}, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"]
    app = [a for a in STATE["apps"] if a["id"] == "app1"][0]
    assert app["policies"] == []


def test_only_the_ticked_dns_records_are_deleted():
    reset_state()
    # two address records on the same hostname
    STATE["dns"].append({"id": "dns-extra", "type": "A", "name": "sonarr.makooy.nl",
                         "content": "192.168.1.77", "proxied": False})
    login()
    page = tc.get("/delete?h=sonarr.makooy.nl&p=").text
    assert "192.168.1.77" in page and "cfargotunnel.com" in page, "both records are shown"
    assert "does not point at the tunnel" in page

    r = post(tc, "/delete", data={"hostname": "sonarr.makooy.nl", "path": "",
                                 "del_dns_ids": "dns1"}, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    left = [d["id"] for d in STATE["dns"] if d["name"] == "sonarr.makooy.nl"]
    assert left == ["dns-extra"], "the record that was not ticked survives"


def test_abandoned_restore_uploads_are_cleaned_up():
    import time as _time
    from app.config import settings as _settings
    stale = _settings.data_dir / "restore-stale.json"
    stale.write_text("{}", encoding="utf-8")
    old = _time.time() - main.RESTORE_TTL - 60
    os.utime(stale, (old, old))
    fresh = _settings.data_dir / "restore-fresh.json"
    fresh.write_text("{}", encoding="utf-8")

    main.purge_restore_uploads()
    assert not stale.exists(), "an hour-old upload is removed"
    assert fresh.exists(), "a recent one is kept"
    fresh.unlink()


def test_delete_everything():
    reset_state()
    login()
    r = post(tc, "/delete", data={"hostname": "sonarr.makooy.nl", "path": "",
                                 "del_route": "1", "del_app": "1", "del_dns_ids": "dns1"},
                follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    assert cf.find_rule(STATE["config"]["ingress"], "sonarr.makooy.nl") is None
    assert not [a for a in STATE["apps"] if a["domain"] == "sonarr.makooy.nl"]
    assert not [d for d in STATE["dns"] if d["name"] == "sonarr.makooy.nl"]


def test_backup_roundtrip():
    reset_state()
    login()
    r = tc.get("/backup/export")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    data = r.json()
    assert data["format"] == backups.FORMAT

    STATE["config"] = cf.remove_rule(STATE["config"], "ha.makooy.nl")
    assert cf.find_rule(STATE["config"]["ingress"], "ha.makooy.nl") is None
    r = post(tc, "/backup/preview", files={"backup": ("b.json", json.dumps(data), "application/json")})
    assert r.status_code == 200 and "ha.makooy.nl" in r.text
    token = r.text.split('name="token" value="')[1].split('"')[0]
    r = post(tc, "/backup/apply", data={"token": token, "restore_routes": "1", "restore_apps": "1"},
                follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    assert cf.find_rule(STATE["config"]["ingress"], "ha.makooy.nl") is not None


def test_non_self_hosted_apps_are_listed_but_not_editable():
    reset_state()
    login()
    page = tc.get("/").text
    assert "workspace.makooy.nl" in page, "a SaaS app is still shown"
    assert "saas" in page and "not editable here" in page

    r = tc.get("/edit?h=workspace.makooy.nl&p=", follow_redirects=False)
    assert r.status_code == 303 and "err=" in r.headers["location"]
    assert "saas+application" in r.headers["location"]

    # a self-hosted one is of course still editable
    r = tc.get("/edit?h=sonarr.makooy.nl&p=")
    assert r.status_code == 200


def test_an_app_covering_several_hostnames_is_matched_on_all_of_them():
    reset_state()
    # a tunnel route for the second hostname of the Books application
    STATE["config"]["ingress"].insert(0, {"hostname": "calibre.makooy.nl",
                                          "service": "http://192.168.1.9:8090"})
    login()
    rows, _ = __import__("asyncio").run(_overview())
    by_host = {r.hostname: r for r in rows}
    assert by_host["calibre.makooy.nl"].app_name == "Books", \
        "matched through destinations, not just domain"
    assert by_host["calibre.makooy.nl"].other_hostnames == ["readarr.makooy.nl"]
    assert len([r for r in rows if r.app_name == "Books"]) == 1, "no duplicate row for one app"


async def _overview():
    main.reset_client()
    services.invalidate_overview()
    return await main.manager().load_overview()


def test_route_order_can_be_changed():
    reset_state()
    login()
    page = tc.get("/routes").text
    assert "ha.makooy.nl" in page and "always last" in page

    before = [r.get("hostname") for r in STATE["config"]["ingress"] if r.get("hostname")]
    r = post(tc, "/routes/move", data={"hostname": before[1], "path": "", "direction": "up"},
             follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    after = [r.get("hostname") for r in STATE["config"]["ingress"] if r.get("hostname")]
    assert after[0] == before[1] and after[1] == before[0], "the two swapped"
    assert cf.is_catch_all(STATE["config"]["ingress"][-1]), "catch-all still last"

    r = post(tc, "/routes/move", data={"hostname": after[0], "path": "", "direction": "up"},
             follow_redirects=False)
    assert "err=" in r.headers["location"], "cannot move the first one up"

    # a route that disappeared between drawing the page and pressing the button
    r = post(tc, "/routes/move", data={"hostname": "gone.makooy.nl", "path": "",
                                       "direction": "up"}, follow_redirects=False)
    assert "no+longer+in+the+tunnel" in r.headers["location"], \
        "must not shuffle a different route when the list changed"


def test_backup_contains_dns_and_can_put_it_back():
    reset_state()
    login()
    data = tc.get("/backup/export").json()
    zones = data["dns_zones"]
    assert zones and zones[0]["zone_name"] == "makooy.nl"
    assert any(r["name"] == "sonarr.makooy.nl" for r in zones[0]["records"])

    # the record disappears, a restore brings it back
    STATE["dns"] = [d for d in STATE["dns"] if d["name"] != "sonarr.makooy.nl"]
    r = post(tc, "/backup/preview",
             files={"backup": ("b.json", json.dumps(data), "application/json")})
    assert r.status_code == 200 and "Missing DNS records" in r.text
    token = r.text.split('name="token" value="')[1].split('"')[0]
    r = post(tc, "/backup/apply", data={"token": token, "restore_dns": "1"},
             follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    assert any(d["name"] == "sonarr.makooy.nl" for d in STATE["dns"])


def test_restore_never_overwrites_an_existing_dns_record():
    reset_state()
    login()
    data = tc.get("/backup/export").json()
    for record in STATE["dns"]:
        if record["name"] == "sonarr.makooy.nl":
            record["content"] = "192.168.1.250"
            record["type"] = "A"
    r = post(tc, "/backup/preview",
             files={"backup": ("b.json", json.dumps(data), "application/json")})
    token = r.text.split('name="token" value="')[1].split('"')[0]
    post(tc, "/backup/apply", data={"token": token, "restore_dns": "1"}, follow_redirects=False)
    kept = [d for d in STATE["dns"] if d["name"] == "sonarr.makooy.nl"]
    assert len(kept) == 1 and kept[0]["content"] == "192.168.1.250", "left exactly as it was"


def test_a_new_app_does_not_accept_one_time_pin_by_default():
    """An empty allowed_idps means Cloudflare accepts every login method, PIN
    included. The form must always send the choice explicitly."""
    reset_state()
    login()
    page = tc.get("/new").text
    assert 'value="idp-google" selected' in page
    assert 'value="idp-otp" selected' not in page, "one-time PIN is not preselected"
    assert "sends a PIN by email" in page
    assert "One-time PIN — onetimepin" in page, "a nameless provider still reads properly"

    r = post(tc, "/new", data={
        "hostname": "vmfleet", "zone": "makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8080",
        "create_app": "1", "app_name": "Vmfleet", "policies": ["pol1"],
        "allowed_idps": ["idp-google"],
    }, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]

    app = [a for a in STATE["apps"] if a["domain"] == "vmfleet.makooy.nl"][0]
    assert app["allowed_idps"] == ["idp-google"], "only Google may be used"
    assert app["auto_redirect_to_identity"] is False


def test_accepting_every_method_has_to_be_asked_for():
    reset_state()
    login()
    base = {"hostname": "open", "zone": "makooy.nl", "scheme": "http",
            "origin_host": "192.168.1.9", "origin_port": "8080",
            "create_app": "1", "app_name": "Open", "policies": ["pol1"]}

    # nothing picked and the box unticked is a mistake, not "allow everything"
    r = post(tc, "/new", data=base, follow_redirects=False)
    assert "err=" in r.headers["location"]
    assert "at+least+one+identity+provider" in r.headers["location"]
    assert not [a for a in STATE["apps"] if a["domain"] == "open.makooy.nl"]

    r = post(tc, "/new", data={**base, "accept_all_idps": "1"}, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    app = [a for a in STATE["apps"] if a["domain"] == "open.makooy.nl"][0]
    assert app["allowed_idps"] == [], "empty is Cloudflare's way of saying all of them"


def test_login_methods_can_be_changed_on_an_existing_app():
    reset_state()
    STATE["apps"][0]["allowed_idps"] = []  # an app that currently accepts everything
    login()
    page = tc.get("/edit?h=sonarr.makooy.nl&p=").text
    assert 'id="accept_all_idps" name="accept_all_idps" value="1"\n           checked' in page \
        or 'name="accept_all_idps" value="1" checked' in page.replace("\n           ", " ")

    r = post(tc, "/edit", data={
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "sonarr.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8989",
        "create_app": "1", "app_name": "Sonarr", "session_duration": "24h",
        "allowed_idps": ["idp-google"],
        "policies": ["pol1"], "allowed_idps": ["idp-google"], "auto_redirect": "1",
    }, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    app = [a for a in STATE["apps"] if a["id"] == "app1"][0]
    assert app["allowed_idps"] == ["idp-google"]
    assert app["auto_redirect_to_identity"] is True
    assert app["custom_deny_message"] == "Ask Marc for access", "the rest still survives"


def test_an_unreadable_provider_list_never_opens_an_app_up(monkeypatch):
    """Without the identity-provider permission the list comes back empty. That
    must not be written back as "accept everything"."""
    reset_state()
    STATE["apps"][0]["allowed_idps"] = ["idp-google"]
    login()

    async def no_providers():
        return []

    monkeypatch.setattr(main, "safe_identity_providers", no_providers)

    page = tc.get("/edit?h=sonarr.makooy.nl&p=").text
    assert "cannot be read" in page and "Identity Providers" in page
    assert 'name="allowed_idps"' not in page, "no empty picker to submit by accident"

    r = post(tc, "/edit", data={
        "original_hostname": "sonarr.makooy.nl", "original_path": "",
        "hostname": "sonarr.makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "8989",
        "create_app": "1", "app_name": "Sonarr", "session_duration": "24h",
        "policies": ["pol1"],
    }, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    app = [a for a in STATE["apps"] if a["id"] == "app1"][0]
    assert app["allowed_idps"] == ["idp-google"], "left exactly as it was"


def test_creating_an_app_is_refused_while_the_list_is_unreadable(monkeypatch):
    reset_state()
    login()

    async def no_providers():
        return []

    monkeypatch.setattr(main, "safe_identity_providers", no_providers)
    base = {"hostname": "blind", "zone": "makooy.nl", "scheme": "http",
            "origin_host": "192.168.1.9", "origin_port": "8080",
            "create_app": "1", "app_name": "Blind", "policies": ["pol1"]}

    r = post(tc, "/new", data=base, follow_redirects=False)
    assert "err=" in r.headers["location"]
    assert "Identity+Providers" in r.headers["location"]
    assert not [a for a in STATE["apps"] if a["domain"] == "blind.makooy.nl"]

    # deliberately accepting everything still works
    r = post(tc, "/new", data={**base, "accept_all_idps": "1"}, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    app = [a for a in STATE["apps"] if a["domain"] == "blind.makooy.nl"][0]
    assert app["allowed_idps"] == []


def test_the_default_login_methods_can_be_set_once():
    reset_state()
    login()
    r = post(tc, "/settings/idps", data={"default_idps": ["idp-otp"]}, follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    assert settings.default_idps == ["idp-otp"]
    page = tc.get("/new").text
    assert 'value="idp-otp" selected' in page and 'value="idp-google" selected' not in page

    post(tc, "/settings/idps", data={}, follow_redirects=False)
    assert settings.default_idps == []
    page = tc.get("/new").text
    assert 'value="idp-google" selected' in page, "back to everything but the PIN"
    assert 'value="idp-otp" selected' not in page


# ------------------------------------------------------- export & backups ---
def test_csv_export():
    reset_state()
    login()
    r = tc.get("/export.csv")
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    body = r.text
    assert body.startswith("\ufeff"), "a BOM so Excel reads UTF-8 properly"
    lines = body.lstrip("\ufeff").strip().splitlines()
    assert lines[0].startswith("Hostname;Path;Origin / service"), lines[0]
    assert any(line.startswith("sonarr.makooy.nl;") for line in lines)
    assert any("GoogleIDP-Marc-Policy" in line for line in lines)


def test_xlsx_export():
    from openpyxl import load_workbook
    reset_state()
    login()
    r = tc.get("/export.xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml")
    sheet = load_workbook(io.BytesIO(r.content)).active
    assert sheet.title == "Applications"
    assert [c.value for c in sheet[1]][:3] == ["Hostname", "Path", "Origin / service"]
    hostnames = [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)]
    assert "sonarr.makooy.nl" in hostnames and "ha.makooy.nl" in hostnames


def test_stored_backup_roundtrip():
    reset_state()
    login()
    for item in backups.listing():
        backups.remove(item["name"])

    r = post(tc, "/backup/create", follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    stored = backups.listing()
    assert len(stored) == 1 and stored[0]["kind"] == "manual"
    assert stored[0]["routes"] == 2 and stored[0]["apps"] == len(STATE["apps"])

    page = tc.get("/backup").text
    assert stored[0]["when"] in page and "by hand" in page

    # restoring a stored backup goes through the same diff screen
    STATE["config"] = cf.remove_rule(STATE["config"], "ha.makooy.nl")
    r = post(tc, "/backup/stored", data={"name": stored[0]["name"]})
    assert r.status_code == 200 and "ha.makooy.nl" in r.text
    token = r.text.split('name="token" value="')[1].split('"')[0]
    r = post(tc, "/backup/apply", data={"token": token, "restore_routes": "1"},
                follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    assert cf.find_rule(STATE["config"]["ingress"], "ha.makooy.nl")

    r = tc.get(f"/backup/download?name={stored[0]['name']}")
    assert r.status_code == 200 and r.json()["format"] == backups.FORMAT

    r = post(tc, "/backup/remove", data={"name": stored[0]["name"]}, follow_redirects=False)
    assert "err=" not in r.headers["location"]
    assert backups.listing() == []


def test_backup_names_cannot_escape_the_backup_directory():
    reset_state()
    login()
    for attempt in ("../secret.key", "..%2Fconfig.json", "/etc/passwd", "config.json"):
        assert backups.resolve(attempt) is None, attempt
        r = tc.get(f"/backup/download?name={attempt}", follow_redirects=False)
        assert r.status_code == 303 and "err=" in r.headers["location"]


def test_backup_is_written_before_a_delete():
    reset_state()
    login()
    for item in backups.listing():
        backups.remove(item["name"])

    r = post(tc, "/delete", data={"hostname": "sonarr.makooy.nl", "path": "", "del_route": "1"},
                follow_redirects=False)
    assert "err=" not in r.headers["location"], r.headers["location"]
    stored = backups.listing()
    assert len(stored) == 1 and stored[0]["kind"] == "predelete"
    assert "sonarr" in stored[0]["name"]
    # the backup still holds the route that was just removed
    data = backups.read(stored[0]["name"])
    assert cf.find_rule(data["tunnel_config"]["ingress"], "sonarr.makooy.nl")
    for item in stored:
        backups.remove(item["name"])


def test_backup_before_delete_can_be_switched_off():
    reset_state()
    login()
    for item in backups.listing():
        backups.remove(item["name"])
    post(tc, "/settings/backups", data={"auto_backup": "1", "auto_backup_hours": "24",
                                       "backup_keep": "14"}, follow_redirects=False)
    assert settings.backup_before_delete is False

    post(tc, "/delete", data={"hostname": "sonarr.makooy.nl", "path": "", "del_route": "1"},
            follow_redirects=False)
    assert backups.listing() == [], "no backup when the option is off"

    post(tc, "/settings/backups", data={"auto_backup": "1", "auto_backup_hours": "24",
                                       "backup_keep": "14", "backup_before_delete": "1"},
            follow_redirects=False)
    assert settings.backup_before_delete is True


def test_a_failed_predelete_backup_asks_before_continuing(monkeypatch):
    reset_state()
    login()
    for item in backups.listing():
        backups.remove(item["name"])

    async def broken(reason):
        raise CloudflareError("Cloudflare is unreachable")

    monkeypatch.setattr(main, "collect_backup", broken)
    data = {"hostname": "sonarr.makooy.nl", "path": "", "del_route": "1", "del_app": "1"}

    r = post(tc, "/delete", data=data)
    assert r.status_code == 200, "a failed backup must not delete anything yet"
    assert "did not work" in r.text and "Cloudflare is unreachable" in r.text
    assert "Delete without a backup" in r.text
    assert cf.find_rule(STATE["config"]["ingress"], "sonarr.makooy.nl"), "route untouched"
    assert [a for a in STATE["apps"] if a["id"] == "app1"], "app untouched"

    r = post(tc, "/delete", data={**data, "confirmed": "1"}, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    assert "backup+failed" in r.headers["location"], "the result says the backup was skipped"
    assert cf.find_rule(STATE["config"]["ingress"], "sonarr.makooy.nl") is None


def test_backups_from_before_the_rename_still_restore():
    reset_state()
    login()
    data = tc.get("/backup/export").json()
    assert data["format"] == "cf0t-manager/1"
    legacy = dict(data, format="cfzerotrustcontrol/1")
    r = post(tc, "/backup/preview",
                files={"backup": ("old.json", json.dumps(legacy), "application/json")})
    assert r.status_code == 200 and "Check this before restoring" in r.text


def test_retention_keeps_only_the_newest():
    import time as _time
    for item in backups.listing():
        backups.remove(item["name"])
    for i in range(5):
        path = backups.write({"format": backups.FORMAT, "n": i}, "auto")
        assert path.exists(), "each write gets its own file, even within one second"
        os.utime(path, (_time.time() - (10 - i), _time.time() - (10 - i)))
    assert len(backups.listing()) == 5
    assert backups.prune("auto", 2) == 3
    left = backups.listing()
    assert len(left) == 2
    assert [backups.read(e["name"])["n"] for e in left] == [4, 3], "the newest two survive"
    for item in left:
        backups.remove(item["name"])


def test_the_global_throttle_cannot_be_sidestepped_by_changing_address():
    """Anyone who can set X-Forwarded-For gets a fresh per-address budget, so a
    global counter has to catch it as well."""
    auth.reset_throttle()
    for i in range(auth.MAX_ATTEMPTS_GLOBAL):
        auth.register_failure(f"10.0.0.{i % 200}")
    assert auth.lockout_remaining("10.0.0.99"), "a brand new address is locked out too"
    auth.reset_throttle()
    assert not auth.lockout_remaining("10.0.0.99")


def test_the_throttle_table_does_not_grow_without_bound():
    auth.reset_throttle()
    for i in range(auth.MAX_TRACKED_ADDRESSES + 200):
        auth.register_failure(f"10.1.{i // 250}.{i % 250}")
    assert len(auth._failures) <= auth.MAX_TRACKED_ADDRESSES
    auth.reset_throttle()


def test_logging_in_hands_out_a_new_csrf_token():
    auth.reset_throttle()
    fresh = TestClient(main.app)
    fresh.get("/login")
    before = fresh.cookies.get(main.CSRF_COOKIE)
    post(fresh, "/login", data={"password": "secret"}, follow_redirects=False)
    after = fresh.cookies.get(main.CSRF_COOKIE)
    assert before and after and before != after, "a token from before the login is not reused"


def test_backup_due_logic():
    import time as _time
    for item in backups.listing():
        backups.remove(item["name"])
    assert main.backup_due() is True, "with no backup at all, one is due"

    backups.write({"format": backups.FORMAT}, "auto")
    assert main.backup_due() is False, "one was just written"
    assert main.backup_due(_time.time() + 25 * 3600) is True, "a day later it is due again"

    settings.save(auto_backup=False)
    assert main.backup_due(_time.time() + 25 * 3600) is False, "unless it is switched off"
    settings.save(auto_backup=True)
    for item in backups.listing():
        backups.remove(item["name"])


def test_changes_are_recorded():
    reset_state()
    login()
    path = settings.data_dir / changelog.FILE
    path.unlink(missing_ok=True)

    post(tc, "/new", data={
        "hostname": "grafana", "zone": "makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "3000",
    }, follow_redirects=False)
    post(tc, "/delete", data={"hostname": "grafana.makooy.nl", "path": "", "del_route": "1"},
         follow_redirects=False)

    entries = changelog.entries()
    assert [e["action"] for e in entries[:2]] == ["deleted", "created"], "newest first"
    assert all(e["target"] == "grafana.makooy.nl" for e in entries[:2])
    assert entries[0]["from"] == "testclient"

    page = tc.get("/changes").text
    assert "grafana.makooy.nl" in page and "deleted" in page
    path.unlink(missing_ok=True)


def test_the_change_log_is_trimmed():
    path = settings.data_dir / changelog.FILE
    path.unlink(missing_ok=True)
    for i in range(changelog.MAX_LINES + 50):
        changelog.record("edited", f"host{i}.makooy.nl")
    assert len(path.read_text().splitlines()) <= changelog.MAX_LINES
    assert changelog.entries(1)[0]["target"].endswith(f"{changelog.MAX_LINES + 49}.makooy.nl")
    path.unlink(missing_ok=True)


# -------------------------------------------------------------- unit bits ---
def test_ingress_helpers():
    cfg = {"ingress": [{"hostname": "a.nl", "service": "http://1:1"}, {"service": "http_status:404"}]}
    out = cf.upsert_rule(cfg, {"hostname": "b.nl", "service": "http://2:2"})
    assert [r.get("hostname") for r in out["ingress"]] == ["a.nl", "b.nl", None]
    out = cf.upsert_rule(out, {"hostname": "a.nl", "service": "http://9:9"})
    assert len(out["ingress"]) == 3 and out["ingress"][0]["service"] == "http://9:9"
    out = cf.remove_rule(out, "a.nl")
    assert len(out["ingress"]) == 2 and cf.is_catch_all(out["ingress"][-1])
    out = cf.upsert_rule({"ingress": []}, {"hostname": "c.nl", "service": "http://3:3"})
    assert out["ingress"][-1]["service"] == "http_status:404", "a config without a catch-all gets one"


def test_service_url_parsing():
    assert services.build_service_url("http", "192.168.1.9", "8096") == "http://192.168.1.9:8096"
    assert services.build_service_url("https", "1.2.3.4", "") == "https://1.2.3.4"
    assert services.build_service_url("http", "", "", "http_status:404") == "http_status:404"
    assert services.parse_service_url("https://192.168.1.18:8006") == {
        "scheme": "https", "host": "192.168.1.18", "port": "8006", "raw": ""}
    assert services.normalize_domain("App.Example.com/api/") == ("app.example.com", "/api")


def test_password_hashing():
    stored = auth.hash_password("hunter2hunter2")
    assert auth.verify_hash("hunter2hunter2", stored)
    assert not auth.verify_hash("something else", stored)
    assert not auth.verify_hash("hunter2hunter2", "garbage")
    assert auth.password_problem("short", "short")
    assert auth.password_problem("longenough", "different")
    assert auth.password_problem("longenough", "longenough") == ""


def test_existing_dns_record_is_never_overwritten_silently():
    reset_state()
    STATE["dns"].append({"id": "dns9", "type": "A", "name": "nas.makooy.nl",
                         "content": "192.168.1.50", "proxied": False})
    login()
    r = post(tc, "/new", data={
        "hostname": "nas", "zone": "makooy.nl", "scheme": "http",
        "origin_host": "192.168.1.9", "origin_port": "5000", "create_dns": "1",
    }, follow_redirects=False)
    assert r.status_code == 200, "a conflict shows a confirmation screen, it does not redirect"
    assert "already has a DNS record" in r.text and "192.168.1.50" in r.text

    record = [d for d in STATE["dns"] if d["name"] == "nas.makooy.nl"][0]
    assert record["type"] == "A" and record["content"] == "192.168.1.50", "untouched"
    assert cf.find_rule(STATE["config"]["ingress"], "nas.makooy.nl"), "route was still created"


def test_dns_record_is_overwritten_once_confirmed():
    reset_state()
    STATE["dns"].append({"id": "dns9", "type": "A", "name": "nas.makooy.nl",
                         "content": "192.168.1.50", "proxied": False})
    login()
    r = post(tc, "/dns/overwrite", data={"hostname": "nas.makooy.nl"}, follow_redirects=False)
    assert r.status_code == 303 and "err=" not in r.headers["location"], r.headers["location"]
    record = [d for d in STATE["dns"] if d["name"] == "nas.makooy.nl"][0]
    assert record["type"] == "CNAME"
    assert record["content"] == f"{TUNNEL}.cfargotunnel.com"
    assert record["proxied"] is True
