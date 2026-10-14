"""CF0T-Manager - web UI for managing Cloudflare Zero Trust apps and tunnel."""
from __future__ import annotations

import csv
import io
import json
import asyncio
import hmac
import logging
import secrets
import time
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any
from urllib.parse import urlencode

from fastapi import Depends, FastAPI, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates

from . import auth, backups, cf, changelog, services
from .cf import CloudflareClient, CloudflareError
from .config import active_tunnel_id, set_active_tunnel_id, settings
from .services import ZeroTrustManager

logging.basicConfig(
    level=logging.INFO,
    format="%(levelname)s:     %(message)s",  # lines up with uvicorn's own output
)
logger = logging.getLogger("cf0t")
BASE_DIR = Path(__file__).resolve().parent
OPEN_PATHS = {"/healthz", "/favicon.ico"}

_client: CloudflareClient | None = None
_client_key: tuple[str, str] | None = None
_stale_clients: list[CloudflareClient] = []


async def collect_backup(reason: str) -> dict:
    """Everything a backup holds, fetched in parallel."""
    mgr = manager()
    tunnel, config, apps, policies = await asyncio.gather(
        client().get_tunnel(mgr.tunnel_id),
        client().get_tunnel_config(mgr.tunnel_id),
        client().list_access_apps(),
        client().list_policies(),
    )
    try:
        dns = await mgr.dns_snapshot(await client().list_zones())
    except CloudflareError as exc:
        logger.warning("backup without DNS records: %s", exc)
        dns = []
    return backups.build_payload(mgr.tunnel_id, tunnel.get("name", ""), config, apps,
                                 policies, reason, dns)


async def store_backup(kind: str, reason: str, label: str = "") -> tuple[str, str]:
    """Write a backup into /data and trim the old ones.
    Returns (file name, error message); never raises."""
    try:
        payload = await collect_backup(reason)
        path = backups.write(payload, kind, label)
        backups.prune(kind, settings.backup_keep)
        return path.name, ""
    except Exception as exc:
        logger.warning("could not store a %s backup: %s", kind, exc)
        return "", str(exc) or exc.__class__.__name__


def backup_due(now: float | None = None) -> bool:
    """Whether an automatic backup should be written right now."""
    if not (settings.auto_backup and settings.configured and active_tunnel_id()):
        return False
    newest = backups.newest_time("auto")
    if newest is None:
        return True
    return ((now or time.time()) - newest) >= settings.auto_backup_hours * 3600


async def backup_scheduler() -> None:
    """Make an automatic backup whenever the newest one is old enough.
    Checking rather than sleeping a full day keeps it correct across restarts."""
    while True:
        try:
            if backup_due():
                name, error = await store_backup("auto", "automatic")
                if name:
                    logger.info("automatic backup written: %s", name)
                else:
                    logger.warning("automatic backup failed: %s", error)
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            logger.warning("automatic backup failed: %s", exc)
        await asyncio.sleep(900)


RESTORE_TTL = 3600  # an abandoned restore upload is worthless after an hour


def purge_restore_uploads() -> None:
    for path in settings.data_dir.glob("restore-*.json"):
        try:
            if time.time() - path.stat().st_mtime > RESTORE_TTL:
                path.unlink()
        except OSError:
            pass


@asynccontextmanager
async def lifespan(app: FastAPI):
    problem = settings.check_data_writable()
    if problem:
        logger.error("the /data volume is not writable (%s) - settings and backups "
                     "cannot be stored. Give the mounted folder to uid 1000.", problem)
    purge_restore_uploads()
    scheduler = asyncio.create_task(backup_scheduler())
    yield
    scheduler.cancel()
    try:
        await scheduler
    except asyncio.CancelledError:
        pass
    for stale in [*_stale_clients, _client]:
        if stale is not None:
            await stale.aclose()


CSRF_COOKIE = "cf0t_csrf"
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


def client() -> CloudflareClient:
    """Rebuilt automatically when the token or account changes."""
    global _client, _client_key
    key = (settings.api_token, settings.account_id)
    if _client is None or _client_key != key:
        _retire(_client)
        _client = CloudflareClient(*key)
        _client_key = key
    return _client


def forget_cached_data() -> None:
    services.invalidate_overview()


def _retire(old: CloudflareClient | None) -> None:
    """Close a replaced client in the background instead of keeping every one
    of them around until the process stops."""
    if old is None:
        return
    try:
        asyncio.get_running_loop().create_task(old.aclose())
    except RuntimeError:  # no loop yet, e.g. at import time
        _stale_clients.append(old)


def reset_client() -> None:
    global _client, _client_key
    _retire(_client)
    _client = None
    _client_key = None
    forget_cached_data()


async def probe_credentials(token: str, account_id: str) -> tuple[str, list[dict]]:
    """Check a token/account pair without touching the cached client.
    Returns (error message, tunnels)."""
    probe = CloudflareClient(token, account_id)
    try:
        await probe.verify_token()
        return "", await probe.list_tunnels()
    except CloudflareError as exc:
        return str(exc), []
    finally:
        await probe.aclose()


def manager() -> ZeroTrustManager:
    return ZeroTrustManager(client(), active_tunnel_id())


DATA_UNWRITABLE = ("Could not write to the /data volume: {error}. "
                   "Nothing was saved. Check that the mounted folder is writable "
                   "by the user the container runs as (uid 1000).")


def client_ip(request: Request) -> str:
    """With --proxy-headers uvicorn has already resolved X-Forwarded-For."""
    return request.client.host if request.client else "unknown"


def save_settings(**values) -> str:
    """Store settings, returning an error message instead of failing silently."""
    try:
        settings.save(**values)
        return ""
    except OSError as exc:
        logger.error("could not persist settings: %s", exc)
        return DATA_UNWRITABLE.format(error=exc)


def redirect(path: str, ok: str = "", err: str = "") -> RedirectResponse:
    params = {k: v for k, v in (("ok", ok), ("err", err)) if v}
    url = f"{path}?{urlencode(params)}" if params else path
    return RedirectResponse(url, status_code=303)


def render(request: Request, template: str, **ctx) -> HTMLResponse:
    ctx.setdefault("ok_msg", request.query_params.get("ok", ""))
    ctx.setdefault("err_msg", request.query_params.get("err", ""))
    ctx.setdefault("settings", settings)
    ctx.setdefault("tunnel_id", active_tunnel_id())

    token = request.cookies.get(CSRF_COOKIE) or secrets.token_hex(16)
    ctx["csrf_token"] = token
    response = templates.TemplateResponse(request, template, ctx)
    response.set_cookie(CSRF_COOKIE, token, httponly=True, samesite="lax",
                        secure=settings.secure_cookie, path="/")
    return response


def rotate_csrf(response) -> None:
    """A new token when the session changes, so one handed out before logging
    in does not stay valid afterwards."""
    response.set_cookie(CSRF_COOKIE, secrets.token_hex(16), httponly=True, samesite="lax",
                        secure=settings.secure_cookie, path="/")


async def csrf_guard(request: Request) -> None:
    """Every POST must echo the cookie back in the form. Blocks another site
    from submitting to this one with your session attached."""
    if request.method != "POST":
        return
    cookie = request.cookies.get(CSRF_COOKIE, "")
    form = await request.form()  # cached, so the route can read it again
    sent = str(form.get("csrf_token") or "")
    if not cookie or not hmac.compare_digest(cookie, sent):
        raise HTTPException(status_code=400,
                            detail="This form expired. Go back, reload the page and try again.")


app = FastAPI(title="CF0T-Manager", lifespan=lifespan, docs_url=None, redoc_url=None,
              openapi_url=None, dependencies=[Depends(csrf_guard)])
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")


@app.middleware("http")
async def guard(request: Request, call_next):
    path = request.url.path
    if path.startswith("/static") or path in OPEN_PATHS:
        return await call_next(request)

    # Nothing configured yet: everything goes to the first-run wizard.
    if settings.needs_setup:
        if path == "/welcome":
            return await call_next(request)
        return RedirectResponse("/welcome", status_code=303)
    if path == "/welcome":
        return RedirectResponse("/", status_code=303)

    if path == "/login":
        return await call_next(request)
    if not auth.is_logged_in(request):
        return RedirectResponse("/login", status_code=303)
    if not active_tunnel_id() and not path.startswith("/settings") and path != "/logout":
        return redirect("/settings", err="Pick a tunnel first.")
    return await call_next(request)


@app.get("/healthz")
async def healthz() -> JSONResponse:
    return JSONResponse({"status": "ok"})


# ----------------------------------------------------------------- login ---
@app.get("/login", response_class=HTMLResponse)
async def login_form(request: Request):
    if not settings.auth_enabled or auth.is_logged_in(request):
        return RedirectResponse("/", status_code=303)
    return render(request, "login.html")


@app.post("/login")
async def login(request: Request, password: str = Form("")):
    who = client_ip(request)
    blocked_for = auth.lockout_remaining(who)
    if blocked_for:
        return redirect("/login", err=f"Too many attempts. Try again in {blocked_for}.")

    if not auth.check_password(password):
        await asyncio.sleep(0.5)  # never time.sleep here: it blocks every request
        left = auth.register_failure(who)
        if left:
            return redirect("/login", err=f"Wrong password. {left} attempts left.")
        return redirect("/login",
                        err=f"Too many attempts. Try again in {auth.lockout_remaining(who)}.")

    auth.register_success(who)
    response = RedirectResponse("/", status_code=303)
    auth.set_session_cookie(response, auth.make_token())
    rotate_csrf(response)
    return response


@app.post("/logout")
async def logout():
    response = RedirectResponse("/login", status_code=303)
    auth.clear_session_cookie(response)
    rotate_csrf(response)
    return response


# --------------------------------------------------- first-run wizard ---
@app.get("/welcome", response_class=HTMLResponse)
async def welcome_form(request: Request):
    return render(request, "welcome.html",
                  need_password=not (settings.auth_enabled or settings.login_disabled),
                  need_cloudflare=not settings.configured, form={})


@app.post("/welcome")
async def welcome_submit(request: Request):
    form = await request.form()
    need_password = not (settings.auth_enabled or settings.login_disabled)
    need_cloudflare = not settings.configured

    api_token = (form.get("api_token") or "").strip() or settings.api_token
    account_id = (form.get("account_id") or "").strip() or settings.account_id
    kept = {"api_token": "", "account_id": account_id}

    if need_password:
        problem = auth.password_problem(form.get("password") or "", form.get("password_confirm") or "")
        if problem:
            return render(request, "welcome.html", need_password=True,
                          need_cloudflare=need_cloudflare, form=kept, err_msg=problem)

    if need_cloudflare and not (api_token and account_id):
        return render(request, "welcome.html", need_password=need_password, need_cloudflare=True,
                      form=kept, err_msg="Both the API token and the account ID are required.")

    error, tunnels = await probe_credentials(api_token, account_id)
    if error:
        return render(request, "welcome.html", need_password=need_password,
                      need_cloudflare=need_cloudflare, form=kept,
                      err_msg=f"Cloudflare rejected these credentials: {error}")

    try:
        settings.save(api_token=api_token, account_id=account_id)
        if need_password:
            settings.set_password_hash(auth.hash_password(form.get("password")))
        if len(tunnels) == 1:
            set_active_tunnel_id(tunnels[0]["id"])
    except OSError as exc:
        return render(request, "welcome.html", need_password=need_password,
                      need_cloudflare=need_cloudflare, form=kept,
                      err_msg=DATA_UNWRITABLE.format(error=exc))
    reset_client()

    response = RedirectResponse("/", status_code=303)
    auth.set_session_cookie(response, auth.make_token())
    return response


# -------------------------------------------------------------- settings ---
@app.get("/settings", response_class=HTMLResponse)
async def settings_page(request: Request):
    tunnels: list[dict] = []
    token_info: dict | None = None
    error = ""
    try:
        token_info = await client().verify_token()
        tunnels = await client().list_tunnels()
        if not active_tunnel_id() and len(tunnels) == 1:
            set_active_tunnel_id(tunnels[0]["id"])
            return redirect("/", ok=f"Selected tunnel '{tunnels[0]['name']}'.")
    except CloudflareError as exc:
        error = str(exc)
    return render(request, "settings.html", tunnels=tunnels, token_info=token_info,
                  cf_error=error, source=settings.source)


@app.post("/settings/cloudflare")
async def settings_cloudflare(request: Request, api_token: str = Form(""),
                              account_id: str = Form("")):
    api_token = api_token.strip() or settings.api_token
    account_id = account_id.strip() or settings.account_id
    if not (api_token and account_id):
        return redirect("/settings", err="Both the API token and the account ID are required.")
    error, tunnels = await probe_credentials(api_token, account_id)
    if error:
        return redirect("/settings", err=f"Cloudflare rejected these credentials: {error}")

    changed_account = account_id != settings.account_id
    error = save_settings(api_token=api_token, account_id=account_id)
    if error:
        return redirect("/settings", err=error)
    if changed_account or not any(t["id"] == active_tunnel_id() for t in tunnels):
        save_settings(tunnel_id=tunnels[0]["id"] if len(tunnels) == 1 else "")
    reset_client()
    return redirect("/settings", ok="Cloudflare credentials saved.")


@app.post("/settings/tunnel")
async def settings_tunnel(request: Request, tunnel_id: str = Form("")):
    if not tunnel_id:
        return redirect("/settings", err="No tunnel selected.")
    error = save_settings(tunnel_id=tunnel_id)
    if error:
        return redirect("/settings", err=error)
    forget_cached_data()
    return redirect("/", ok="Tunnel saved.")


@app.post("/settings/password")
async def settings_password(request: Request, current_password: str = Form(""),
                            password: str = Form(""), password_confirm: str = Form("")):
    if settings.env_password:
        return redirect("/settings", err="The password is set through CF0T_PASSWORD; change it there.")
    if settings.auth_enabled and not auth.check_password(current_password):
        return redirect("/settings", err="The current password is not correct.")
    problem = auth.password_problem(password, password_confirm)
    if problem:
        return redirect("/settings", err=problem)

    try:
        settings.set_password_hash(auth.hash_password(password))
    except OSError as exc:
        return redirect("/settings", err=DATA_UNWRITABLE.format(error=exc))
    response = redirect("/settings", ok="Password changed. Other sessions have been signed out.")
    auth.set_session_cookie(response, auth.make_token())
    return response


@app.post("/settings/backups")
async def settings_backups(request: Request, auto_backup: str = Form(""),
                           auto_backup_hours: str = Form("24"), backup_keep: str = Form("14"),
                           backup_before_delete: str = Form("")):
    try:
        hours = max(1, min(int(auto_backup_hours), 24 * 30))
        keep = max(1, min(int(backup_keep), 365))
    except ValueError:
        return redirect("/settings", err="Interval and retention must be numbers.")
    error = save_settings(auto_backup=bool(auto_backup), auto_backup_hours=hours,
                          backup_keep=keep, backup_before_delete=bool(backup_before_delete))
    return redirect("/settings", err=error) if error else \
        redirect("/settings", ok="Backup settings saved.")


@app.post("/settings/preferences")
async def settings_preferences(request: Request, session_hours: str = Form("12"),
                               secure_cookie: str = Form(""), default_origin_host: str = Form("")):
    try:
        hours = max(1, min(int(session_hours), 24 * 365))
    except ValueError:
        return redirect("/settings", err="Session duration must be a number of hours.")
    error = save_settings(session_hours=hours, secure_cookie=bool(secure_cookie),
                          default_origin_host=default_origin_host.strip() or "192.168.1.9")
    return redirect("/settings", err=error) if error else \
        redirect("/settings", ok="Preferences saved.")


# ------------------------------------------------------------- dashboard ---
@app.get("/", response_class=HTMLResponse)
async def dashboard(request: Request):
    mgr = manager()
    refresh = request.query_params.get("refresh") == "1"
    try:
        (rows, meta), tunnel = await asyncio.gather(
            mgr.load_overview(refresh=refresh),
            client().get_tunnel(mgr.tunnel_id),
        )
    except CloudflareError as exc:
        return render(request, "index.html", rows=[], meta={}, tunnel=None, err_msg=str(exc))
    return render(request, "index.html", rows=rows, meta=meta, tunnel=tunnel)


def destructive_warnings(form, row=None) -> list[str]:
    """Changes that quietly remove protection, spelled out before they happen."""
    warnings: list[str] = []
    keeps_app = bool(form.get("create_app"))
    policies = form_list(form, "policies")

    if row is not None and row.has_app and not keeps_app:
        warnings.append(
            f"The Access application \u201c{row.app_name}\u201d will be deleted. "
            f"{row.hostname} stays reachable through the tunnel, but without any "
            f"authentication in front of it."
        )
    if keeps_app and not policies:
        had = len(row.policies) if row is not None else 0
        if had:
            warnings.append(
                f"All {had} policy/policies will be removed from this application. "
                f"Cloudflare Access denies everyone while an application has no policy."
            )
        else:
            warnings.append(
                "This application will be created without a policy. Cloudflare Access "
                "denies everyone until you attach one."
            )
    return warnings


def confirmation_page(request: Request, action: str, form, warnings: list[str],
                      confirm_label: str = "Yes, apply it") -> HTMLResponse:
    fields = [(k, v) for k, v in form.multi_items()
              if k not in {"confirmed", "csrf_token"}]
    return render(request, "confirm.html", action=action, fields=fields, warnings=warnings,
                  confirm_label=confirm_label)


EXPORT_COLUMNS = ("Hostname", "Path", "Origin / service", "Access app", "Policies",
                  "Session duration", "In App Launcher", "Tunnel route", "Access app present",
                  "DNS type", "DNS target")


def export_rows(rows) -> list[list[str]]:
    table = []
    for row in rows:
        app = row.app or {}
        table.append([
            row.hostname,
            row.path,
            row.service,
            row.app_name,
            ", ".join(p.get("name", "") for p in row.policies),
            app.get("session_duration", ""),
            "yes" if app.get("app_launcher_visible") else "no" if app else "",
            "yes" if row.has_route else "no",
            "yes" if row.has_app else "no",
            (row.dns or {}).get("type", ""),
            (row.dns or {}).get("content", ""),
        ])
    return table


@app.get("/export.csv")
async def export_csv():
    try:
        rows, _ = await manager().load_overview()
    except CloudflareError as exc:
        return redirect("/", err=str(exc))
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";")  # opens straight into Excel on a Dutch locale
    writer.writerow(EXPORT_COLUMNS)
    writer.writerows(export_rows(rows))
    stamp = time.strftime("%Y%m%d-%H%M%S")
    return Response(
        "\ufeff" + buffer.getvalue(),  # BOM so Excel reads the UTF-8 correctly
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="applications-{stamp}.csv"'},
    )


@app.get("/export.xlsx")
async def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    try:
        rows, _ = await manager().load_overview()
    except CloudflareError as exc:
        return redirect("/", err=str(exc))

    table = export_rows(rows)

    book = Workbook()
    sheet = book.active
    sheet.title = "Applications"
    sheet.append(list(EXPORT_COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for line in table:
        sheet.append(line)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{sheet.max_row}"
    for index, header in enumerate(EXPORT_COLUMNS, start=1):
        longest = max([len(header)] + [len(str(line[index - 1])) for line in table])
        sheet.column_dimensions[get_column_letter(index)].width = min(48, longest + 3)

    stream = io.BytesIO()
    book.save(stream)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    return Response(
        stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="applications-{stamp}.xlsx"'},
    )


# ------------------------------------------------------ add application ---
@app.get("/new", response_class=HTMLResponse)
async def new_form(request: Request):
    try:
        policies, zones = await asyncio.gather(
            client().list_policies(), client().list_zones()
        )
    except CloudflareError as exc:
        return redirect("/", err=str(exc))
    return render(request, "form.html", mode="new", policies=policies, zones=zones,
                  row=None, form={}, selected_policies=[])


@app.post("/new")
async def new_submit(request: Request):
    form = await request.form()
    mgr = manager()
    try:
        hostname, service, path, origin = _parse_route_form(form)
    except ValueError as exc:
        return redirect("/new", err=str(exc))

    warnings = destructive_warnings(form)
    if warnings and not form.get("confirmed"):
        return confirmation_page(request, "/new", form, warnings)

    steps: list[str] = []
    try:
        await mgr.save_route(hostname, path, service, origin)
        steps.append("tunnel route")

        if form.get("create_app"):
            policy_ids = form_list(form, "policies")
            domain = hostname + path if path else hostname
            await mgr.create_app(
                name=(form.get("app_name") or hostname.split(".")[0]).strip(),
                domain=domain,
                policy_ids=policy_ids,
                session_duration=(form.get("session_duration") or "24h"),
                launcher_visible=bool(form.get("launcher_visible")),
            )
            steps.append(f"Access app ({len(policy_ids)} policy/policies)")

        if form.get("create_dns"):
            zones = await client().list_zones()
            steps.append(await mgr.ensure_dns(hostname, zones))
    except services.DnsConflict as conflict:
        return render(request, "dns_conflict.html", conflict=conflict, steps=steps)
    except CloudflareError as exc:
        done = ", ".join(steps) or "nothing"
        return redirect("/", err=f"{exc} — already done: {done}.")

    changelog.record("created", hostname + path, ", ".join(steps), client_ip(request))
    return redirect("/", ok=f"Created {hostname}: " + ", ".join(steps))


# ----------------------------------------------------- edit application ---
@app.get("/edit", response_class=HTMLResponse)
async def edit_form(request: Request):
    hostname = request.query_params.get("h", "")
    path = services.normalize_path(request.query_params.get("p", ""))
    mgr = manager()
    try:
        row, policies, zones = await asyncio.gather(
            mgr.get_row(hostname, path), client().list_policies(), client().list_zones()
        )
    except CloudflareError as exc:
        return redirect("/", err=str(exc))
    if row is None:
        return redirect("/", err=f"{hostname} not found.")
    if not row.editable:
        return redirect("/", err=f"{row.hostname} is a {row.app_type.replace('_', ' ')} "
                               f"application with no tunnel route, so there is nothing to "
                               f"edit here. Change it in the Cloudflare dashboard.")

    origin = (row.rule or {}).get("originRequest") or {}
    parsed = services.parse_service_url(row.service)
    # Show the subdomain on its own when the domain dropdown already carries
    # the zone, so the two fields do not repeat each other.
    subdomain, zone_name = services.split_hostname(row.hostname, zones)
    form: dict[str, Any] = {
        "hostname": subdomain,
        "zone": zone_name,
        "path": row.path,
        "app_name": row.app_name or row.hostname.split(".")[0],
        "session_duration": (row.app or {}).get("session_duration") or "24h",
        "launcher_visible": (row.app or {}).get("app_launcher_visible", True),
        "create_app": row.has_app,
        **parsed,
        **origin,
    }
    return render(request, "form.html", mode="edit", policies=policies, zones=zones, row=row,
                  form=form, selected_policies=[p.get("id") for p in row.policies])


@app.post("/edit")
async def edit_submit(request: Request):
    form = await request.form()
    original_host = (form.get("original_hostname") or "").strip().lower()
    original_path = services.normalize_path(form.get("original_path") or "")
    mgr = manager()

    try:
        hostname, service, path, origin = _parse_route_form(form)
        row = await mgr.get_row(original_host, original_path)
    except (ValueError, CloudflareError) as exc:
        return redirect(f"/edit?h={original_host}&p={original_path}", err=str(exc))
    if row is None:
        return redirect("/", err=f"{original_host} not found.")

    warnings = destructive_warnings(form, row)
    if warnings and not form.get("confirmed"):
        return confirmation_page(request, "/edit", form, warnings)

    steps: list[str] = []
    try:
        if row.has_route:
            await mgr.save_route(hostname, path, service, origin, original_host, original_path)
            steps.append("route updated")
        elif form.get("create_route"):
            await mgr.save_route(hostname, path, service, origin)
            steps.append("route created")

        policy_ids = form_list(form, "policies")
        domain = hostname + path if path else hostname
        name = (form.get("app_name") or hostname.split(".")[0]).strip()
        duration = form.get("session_duration") or "24h"
        visible = bool(form.get("launcher_visible"))

        if row.has_app and form.get("create_app"):
            await mgr.update_app(row.app, name, domain, policy_ids, duration, visible)
            steps.append("Access app updated")
        elif row.has_app and not form.get("create_app"):
            await client().delete_access_app(row.app["id"])
            services.invalidate_overview(mgr.tunnel_id)
            steps.append("Access app deleted")
        elif not row.has_app and form.get("create_app"):
            await mgr.create_app(name, domain, policy_ids, duration, visible)
            steps.append("Access app created")

        # A rename needs its own DNS record. If the old hostname had a dedicated
        # one, carry it over even when the checkbox was left unticked: that
        # checkbox described the old hostname, not this one.
        renamed = hostname != original_host
        inherits_dns = renamed and row.dns_ok
        if form.get("create_dns") or inherits_dns:
            zones = await client().list_zones()
            steps.append(await mgr.ensure_dns(hostname, zones))
            if inherits_dns:
                steps.append(f"{original_host} still has a DNS record of its own, "
                             f"listed on the overview so you can remove it")
    except services.DnsConflict as conflict:
        return render(request, "dns_conflict.html", conflict=conflict, steps=steps)
    except CloudflareError as exc:
        done = ", ".join(steps) or "nothing"
        return redirect("/", err=f"{exc} — already done: {done}.")

    if steps:
        changelog.record("edited", hostname + path, ", ".join(steps), client_ip(request))
    return redirect("/", ok=f"{hostname}: " + ", ".join(steps or ["no changes"]))


# --------------------------------------------------- delete application ---
@app.get("/delete", response_class=HTMLResponse)
async def delete_form(request: Request):
    hostname = request.query_params.get("h", "")
    path = services.normalize_path(request.query_params.get("p", ""))
    mgr = manager()
    try:
        row = await mgr.get_row(hostname, path)
        if row is None:
            return redirect("/", err=f"{hostname} not found.")
        zones = await client().list_zones()
        dns_records = await mgr.dns_records_for(row.hostname, zones)
    except CloudflareError:
        dns_records = []
    return render(request, "delete.html", row=row, dns_records=dns_records)


@app.post("/delete")
async def delete_submit(request: Request):
    form = await request.form()
    hostname = (form.get("hostname") or "").strip().lower()
    path = services.normalize_path(form.get("path") or "")
    mgr = manager()
    try:
        row = await mgr.get_row(hostname, path)
    except CloudflareError as exc:
        return redirect("/", err=str(exc))
    if row is None:
        return redirect("/", err=f"{hostname} not found.")

    steps: list[str] = []
    try:
        if settings.backup_before_delete:
            name, error = await store_backup("predelete", f"before deleting {hostname}", hostname)
            if error and not form.get("confirmed"):
                return confirmation_page(request, "/delete", form, [
                    f"The backup that is normally made before a delete did not work: {error}",
                    f"If you continue, {hostname}{path} is deleted without one, and there is "
                    f"nothing to restore from afterwards.",
                ], confirm_label="Delete without a backup")
            steps.append(f"backup stored ({name})" if name
                         else "backup failed, continued on request")

        if form.get("del_app") and row.has_app:
            await client().delete_access_app(row.app["id"])
            services.invalidate_overview(mgr.tunnel_id)
            steps.append("Access app deleted")
        if form.get("del_route") and row.has_route:
            await mgr.delete_route(hostname, path)
            steps.append("tunnel route deleted")
        record_ids = form_list(form, "del_dns_ids")
        if record_ids:
            zones = await client().list_zones()
            steps.append(await mgr.delete_dns(hostname, zones, record_ids))
    except CloudflareError as exc:
        done = ", ".join(steps) or "nothing"
        return redirect("/", err=f"{exc} — already done: {done}.")

    if steps:
        changelog.record("deleted", hostname + path, ", ".join(steps), client_ip(request))
    return redirect("/", ok=f"{hostname}: " + ", ".join(steps or ["nothing selected"]))


@app.post("/dns/overwrite")
async def dns_overwrite(request: Request, hostname: str = Form("")):
    """Only reached after the user confirmed it on the conflict screen."""
    try:
        zones = await client().list_zones()
        message = await manager().ensure_dns(hostname, zones, overwrite=True)
    except (CloudflareError, services.DnsConflict) as exc:
        return redirect("/", err=str(exc))
    changelog.record("dns overwritten", hostname, message, client_ip(request))
    return redirect("/", ok=f"{hostname}: {message.lower()}")


# ----------------------------------------------------------- route order ---
@app.get("/routes", response_class=HTMLResponse)
async def routes_order(request: Request):
    try:
        config = await client().get_tunnel_config(manager().tunnel_id)
    except CloudflareError as exc:
        return redirect("/", err=str(exc))
    rules, catch_all = cf.split_ingress(config.get("ingress") or [])
    return render(request, "routes.html", rules=rules, catch_all=catch_all)


@app.post("/routes/move")
async def routes_move(request: Request, hostname: str = Form(""), path: str = Form(""),
                      direction: str = Form("")):
    """Identified by hostname and path rather than by position: the list may
    have changed since the page was drawn."""
    mgr = manager()
    try:
        config = await client().get_tunnel_config(mgr.tunnel_id)
        rules, _ = cf.split_ingress(config.get("ingress") or [])
        order = [cf.rule_key(r.get("hostname"), r.get("path")) for r in rules]
        wanted = cf.rule_key(hostname, path)
        if wanted not in order:
            return redirect("/routes", err=f"{hostname} is no longer in the tunnel "
                                           f"configuration. The list has been reloaded.")

        position = order.index(wanted)
        target = position - 1 if direction == "up" else position + 1
        if not 0 <= target < len(order):
            return redirect("/routes", err="That route is already at the end.")

        order[position], order[target] = order[target], order[position]
        await client().put_tunnel_config(mgr.tunnel_id, cf.reorder_rules(config, order))
        services.invalidate_overview(mgr.tunnel_id)
    except CloudflareError as exc:
        return redirect("/routes", err=str(exc))

    changelog.record("reordered", hostname + path, f"moved {direction}", client_ip(request))
    return redirect("/routes", ok=f"Moved {hostname}{path} {direction}.")


@app.get("/changes", response_class=HTMLResponse)
async def changes_page(request: Request):
    return render(request, "changes.html", entries=changelog.entries())


# ---------------------------------------------------------------- backup ---
@app.get("/backup", response_class=HTMLResponse)
async def backup_page(request: Request):
    return render(request, "backup.html", preview=None, stored=backups.listing())


@app.get("/backup/export")
async def backup_export():
    try:
        payload = await collect_backup("manual download")
    except CloudflareError as exc:
        return redirect("/backup", err=str(exc))
    stamp = time.strftime("%Y%m%d-%H%M%S")
    return Response(
        json.dumps(payload, indent=2, ensure_ascii=False),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="cf0t-backup-{stamp}.json"'},
    )


@app.post("/backup/create")
async def backup_create():
    name, error = await store_backup("manual", "made by hand")
    if not name:
        return redirect("/backup", err=f"Nothing was stored: {error}")
    return redirect("/backup", ok=f"Backup stored as {name}.")


@app.get("/backup/download")
async def backup_download(request: Request):
    name = request.query_params.get("name", "")
    path = backups.resolve(name)
    if path is None:
        return redirect("/backup", err="That backup no longer exists.")
    return Response(
        path.read_text(encoding="utf-8"),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{path.name}"'},
    )


@app.post("/backup/remove")
async def backup_remove(name: str = Form("")):
    if not backups.remove(name):
        return redirect("/backup", err="That backup no longer exists.")
    return redirect("/backup", ok=f"Deleted {name}.")


@app.post("/backup/stored", response_class=HTMLResponse)
async def backup_preview_stored(request: Request, name: str = Form("")):
    data = backups.read(name)
    if data is None:
        return redirect("/backup", err="That backup could not be read.")
    return await render_restore_preview(request, data)


@app.post("/backup/preview", response_class=HTMLResponse)
async def backup_preview(request: Request, backup: UploadFile | None = None):
    if backup is None or not backup.filename:
        return redirect("/backup", err="No file selected.")
    try:
        data = json.loads((await backup.read()).decode("utf-8"))
    except (ValueError, UnicodeDecodeError) as exc:
        return redirect("/backup", err=f"Could not read that file: {exc}")
    return await render_restore_preview(request, data)


async def render_restore_preview(request: Request, data: dict):
    """Show what a restore would change, from an upload or a stored backup."""
    if not backups.is_ours(data.get("format") or ""):
        return redirect("/backup", err="That doesn't look like a CF0T-Manager backup.")

    try:
        current_config = await client().get_tunnel_config(manager().tunnel_id)
        current_apps = await client().list_access_apps()
    except CloudflareError as exc:
        return redirect("/backup", err=str(exc))

    backup_rules, _ = cf.split_ingress(data.get("tunnel_config", {}).get("ingress") or [])
    current_keys = {cf.rule_key(r.get("hostname"), r.get("path"))
                    for r in cf.split_ingress(current_config.get("ingress") or [])[0]}
    backup_keys = {cf.rule_key(r.get("hostname"), r.get("path")) for r in backup_rules}

    app_domains = {services.normalize_domain(a.get("domain") or "") for a in current_apps}
    missing_apps = [
        a for a in data.get("access_apps") or []
        if services.normalize_domain(a.get("domain") or "") not in app_domains
        and (a.get("type") or "") == "self_hosted"
    ]

    purge_restore_uploads()
    token = secrets.token_urlsafe(16)
    (settings.data_dir / f"restore-{token}.json").write_text(
        json.dumps(data), encoding="utf-8"
    )

    missing_dns = 0
    try:
        for zone in data.get("dns_zones") or []:
            if not zone.get("zone_id"):
                continue
            present = {(r.get("name") or "").lower()
                       for r in await client().list_dns_records(zone["zone_id"])}
            missing_dns += sum(1 for r in zone.get("records") or []
                               if (r.get("name") or "").lower() not in present)
    except CloudflareError:
        missing_dns = 0

    preview = {
        "token": token,
        "missing_dns": missing_dns,
        "tunnel_name": (data.get("tunnel") or {}).get("name"),
        "exported_at": data.get("exported_at"),
        "routes_total": len(backup_rules),
        "routes_new": sorted(h for h, _ in backup_keys - current_keys),
        "routes_removed": sorted(h for h, _ in current_keys - backup_keys),
        "missing_apps": missing_apps,
    }
    return render(request, "backup.html", preview=preview, stored=backups.listing())


@app.post("/backup/apply")
async def backup_apply(request: Request, token: str = Form(""), restore_routes: str = Form(""),
                       restore_apps: str = Form(""), restore_dns: str = Form("")):
    path = settings.data_dir / f"restore-{Path(token).name}.json"
    if not token or not path.exists():
        return redirect("/backup", err="Upload expired, please try again.")
    data = json.loads(path.read_text(encoding="utf-8"))
    mgr = manager()
    steps: list[str] = []
    try:
        if restore_routes:
            config = data.get("tunnel_config") or {}
            rules, catch_all = cf.split_ingress(config.get("ingress") or [])
            config["ingress"] = cf.build_ingress(rules, catch_all)
            await client().put_tunnel_config(mgr.tunnel_id, config)
            services.invalidate_overview(mgr.tunnel_id)
            steps.append(f"restored {len(rules)} tunnel routes")
        if restore_apps:
            current = {services.normalize_domain(a.get("domain") or "")
                       for a in await client().list_access_apps()}
            created = 0
            for app_def in data.get("access_apps") or []:
                domain = app_def.get("domain") or ""
                if (app_def.get("type") or "") != "self_hosted":
                    continue
                if services.normalize_domain(domain) in current:
                    continue
                await mgr.recreate_app(app_def)
                created += 1
            steps.append(f"created {created} Access app(s)")
        if restore_dns:
            created, skipped, unreachable = await mgr.restore_dns(data.get("dns_zones") or [])
            steps.append(f"created {created} DNS record(s)"
                         + (f", left {skipped} existing one(s) alone" if skipped else "")
                         + (f", could not reach {', '.join(unreachable)}" if unreachable else ""))
    except CloudflareError as exc:
        return redirect("/backup", err=f"{exc} — already done: {', '.join(steps) or 'nothing'}.")
    finally:
        path.unlink(missing_ok=True)
    changelog.record("restored", data.get("tunnel", {}).get("name", "tunnel"),
                     ", ".join(steps), client_ip(request))
    return redirect("/", ok="Restore finished: " + ", ".join(steps or ["nothing selected"]))


# -------------------------------------------------------------- helpers ----
def form_list(form, key: str) -> list[str]:
    """Multi-select fields: FormData.getlist, otherwise a plain dict."""
    if hasattr(form, "getlist"):
        return [v for v in form.getlist(key) if v]
    value = form.get(key)
    if isinstance(value, list):
        return [v for v in value if v]
    return [value] if value else []


def _parse_route_form(form) -> tuple[str, str, str, dict]:
    hostname = (form.get("hostname") or "").strip().lower().rstrip(".")
    zone = (form.get("zone") or "").strip().lower()
    if zone and not hostname.endswith(zone):
        hostname = f"{hostname}.{zone}" if hostname else zone
    if not hostname or "." not in hostname:
        raise ValueError("Enter a valid hostname.")
    service = services.build_service_url(
        form.get("scheme", "http"), form.get("origin_host", ""),
        form.get("origin_port", ""), form.get("raw_service", ""),
    )
    path = services.normalize_path(form.get("path") or "")
    origin = services.build_origin_request(form)
    return hostname, service, path, origin
